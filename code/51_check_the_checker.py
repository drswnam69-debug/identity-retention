#!/usr/bin/env python3
"""Mutation test for 50_consistency_check.py.

A checker that only ever passes is worthless. This injects known defects into
copies of the manuscript and confirms each one is caught.
"""
import os, re, shutil, subprocess, sys, tempfile


import os as _os, sys as _sys
import os as _os, sys as _sys
_here = _os.path.dirname(_os.path.abspath(__file__))
_cands = [_here, _os.path.join(_here, "rsi", "code"), _os.path.join(_here, "code"),
          _os.path.join(_os.path.dirname(_here), "code")]
if _os.environ.get("IR_ROOT"):
    _cands.insert(0, _os.path.join(_os.environ["IR_ROOT"], "code"))
for _c in _cands:
    if _os.path.exists(_os.path.join(_c, "paths.py")):
        if _c not in _sys.path:
            _sys.path.insert(0, _c)
        break
from paths import ROOT as IR_ROOT, RESULTS as IR_RESULTS, GENESETS as IR_GENESETS, \
    DATA as IR_DATA, CODE as IR_CODE, FIGURES as IR_FIGURES, DOCS as IR_DOCS
_here = _os.path.dirname(_os.path.abspath(__file__))
for _c in (_here, _os.path.join(_here, "rsi", "code"), _os.path.join(_here, "code"),
           _os.path.join(_os.path.dirname(_here), "code")):
    if _os.path.exists(_os.path.join(_c, "paths.py")):
        if _c not in _sys.path:
            _sys.path.insert(0, _c)
        break
from paths import ROOT as IR_ROOT, RESULTS as IR_RESULTS, CODE as IR_CODE, DOCS as IR_DOCS

HOME = IR_DOCS
SRC = f"{HOME}/manuscript_v4.md"
CHK = f"{HOME}/rsi/code/50_consistency_check.py"

MUTATIONS = [
    # a dataset citation pointing at the wrong entry: every number stays valid
    # and every accession is still named, only the pairing is wrong
    ("a dataset accession cited to the wrong reference entry",
     "GSE167523 [36]", "GSE167523 [35]"),
    ("a dataset reference stripped of its persistent link",
     " PXD006512. https://www.ebi.ac.uk/pride/archive/projects/PXD006512",
     " PXD006512."),
    # companion-file defects: the checker reads these files too
    ("alt text describing the wrong figure",
     "**Figure 9.** Four panels comparing an array cohort",
     "**Figure 8.** Four panels comparing an array cohort",
     f"{IR_DOCS}/Figure_alt_text.md"),
    ("the cover letter claiming a benchmark the manuscript declines",
     "The premise check and the ratio-stability guard live inside that function",
     "Adjusting for identity is not what adjusting for tumor purity does. "
     "The premise check and the ratio-stability guard live inside that function",
     f"{IR_DOCS}/CoverLetter_GigaScience.md"),
    # this sentence now lives in the supplementary note, which is part of the
    # same checked corpus; the mutation follows it there
    ("a literature-sourced percentage silently changed",
     "cirrhosis was recorded in **54%** of the 115 patients",
     "cirrhosis was recorded in **58%** of the 115 patients",
     f"{IR_DOCS}/SupplementaryNote_v1.md"),
    ("the non-replication dropped from the abstract",
     "nor did one arm of the example replicate on RNA sequencing; a",
     "; a"),
    ("an amendment that fixes an outcome dropped from the manuscript",
     "\u00a76ab", "that amendment"),
    ("the letter's word count left stale",
     "about 23,200 words from Background", "about 15,000 words from Background",
     f"{IR_DOCS}/CoverLetter_GigaScience.md"),
    ("a figure script using a glyph the font lacks",
     'r"$9.1 \\times 10^{-10}$"', '"9.1 \\u00d7 10\\u207b\\u00b9\\u2070"',
     f"{IR_DOCS}/make_figure5.py"),
    ("a British spelling the old word list missed",
     "the dotted gray lines mark", "the dotted grey lines mark"),
    # the journal's own heading is "Acknowledgements"; the exemption is for that
    # heading alone, so a British spelling anywhere else must still be caught
    ("a British spelling in the Acknowledgements text",
     "**Acknowledgements.** None.",
     "**Acknowledgements.** The figures were colour-checked by the author."),
    ("Table 1 calling lung the second tissue",
     "Lung adenocarcinoma; 58 pairs | Third tissue (\u00a76y)",
     "Lung adenocarcinoma; 58 pairs | Second tissue (\u00a76x)"),
    ("a gene set attributed to the protocol that the protocol does not name",
     "I substituted the two directional Gene Ontology sets under the clause of "
     "\u00a76d that requires the set actually used to be recorded with the result, "
     "and I record here that the substitution is mine and was not pre-specified.",
     "These are the two sets \u00a76d named.",
     f"{IR_DOCS}/SupplementaryNote_v1.md"),
    ("a British spelling left in the cover letter",
     "A tumor specimen and the adjacent tissue",
     "A tumour specimen and the adjacent tissue",
     f"{IR_DOCS}/CoverLetter_GigaScience.md"),
    ("the cover letter left disagreeing with the manuscript's amendment count",
     "Five amendments in all are not pre-registered",
     "Six amendments in all are not pre-registered"),
    ("the submission metadata file left stale after an abstract revision",
     "it ranged 0.004 to 2.345, median 0.758",
     "it ranged 0.004 to 2.345, median 0.812"),
    # co-occurrence defects: each number is real, only the pairing is wrong
    ("an interval bound taken from a sibling model",
     "(−0.248, −0.390 to −0.106, *P* = 6.8 × 10⁻⁴)",
     "(−0.248, −0.390 to −0.109, *P* = 6.8 × 10⁻⁴)"),
    ("two genes' P values swapped",
     "*CYB5R3* fell (−0.387, *P* = 3.0 × 10⁻⁴) while *AIFM2* rose (+0.308, *P* = 4.9 × 10⁻³)",
     "*CYB5R3* fell (−0.387, *P* = 4.9 × 10⁻³) while *AIFM2* rose (+0.308, *P* = 3.0 × 10⁻⁴)"),
    ("a correlation paired with another correlation's P",
     "ρ = −0.261, *P* = 1.2 × 10⁻⁴", "ρ = −0.261, *P* = 1.2 × 10⁻¹³"),
    ("an estimate replaced by another dataset's estimate",
     "DRAIN fell strongly (−0.944, *P* = 2.3 × 10⁻²¹)",
     "DRAIN fell strongly (−0.716, *P* = 2.3 × 10⁻²¹)"),
    ("a table cell's interval bound moved",
     "+0.373 (0.031 to 0.715), *P* = 0.033",
     "+0.373 (0.031 to 0.735), *P* = 0.033"),
    ("a number changed to one nothing supports",
     "median of **0.758**", "median of **0.799**"),
    ("a retention percentage altered",
     "REDUCTION retained **89%**", "REDUCTION retained **77%**"),
    ("a P value altered",
     "95% CI 0.614 to 0.913, *P* = 1.0 × 10⁻¹⁹", "95% CI 0.614 to 0.913, *P* = 1.0 × 10⁻¹⁴"),
    ("a sample count altered",
     "TCGA-LUAD, 589 samples giving 58 pairs", "TCGA-LUAD, 578 samples giving 58 pairs"),
    ("an interval width altered",
     "Median interval widths are 0.327", "Median interval widths are 0.427"),
    ("a per-gene table value altered",
     "| *MTARC2* | −0.923 (1.2 × 10⁻²⁴)", "| *MTARC2* | −0.823 (1.2 × 10⁻²⁴)"),
    ("a figure legend pushed over 300 words",
     "**Figure 12. What the source proteogenomic study settles (§6aa).**",
     "**Figure 12. What the source proteogenomic study settles (§6aa).** " + "filler word " * 160),
    ("a correlation altered",
     "ρ = +0.548", "ρ = +0.648"),
    ("an em dash introduced",
     "The two arms respond to it differently", "The two arms respond — differently"),
    ("a British spelling introduced",
     "hematoxylin and eosin staining", "haematoxylin and eosin staining"),
    ("first person plural introduced",
     "I defined and locked the identity-retention fraction",
     "We defined and locked the identity-retention fraction"),
    ("a citation to a reference that does not exist",
     "collection of the Molecular Signatures Database (MSigDB) [", "collection of the Molecular Signatures Database (MSigDB) [99"),
    ("an amendment cited that the protocol lacks",
     "Under §6z, written before any survival value", "Under §6zz, written before any survival value"),
    ("the amendment count made wrong",
     "Thirty-one amendments (§6a to §6ae)", "Thirty amendments (§6a to §6ae)"),
    ("the lock digest altered",
     "`7a2bf934…b680046`", "`7a2bf934…b680047`"),
    ("a table title made too long",
     "**Table 1. Nine public transcriptome series and two published proteomes, with role and pairing.**",
     "**Table 1. Nine public transcriptome series and two published proteomes, each identified here by its accession or its project name, with role and pairing.**"),
    ("the abstract pushed over the word limit",
     "**Conclusions.** How much", "**Conclusions.** " + "padding word " * 40 + "How much"),
]


# Two defects cannot be written as a substitution in a text file, because the
# things they damage are not text: the workbook the editor receives, and the
# set of files in the package. Each is a pair of functions, applied and then
# undone, and each is counted with the substitutions above.
PACKAGE = f"{HOME}/submission_GigaScience"
SM2_XLSX = f"{PACKAGE}/04_Supplementary Material 2_Supplementary tables.xlsx"
_STRAY_ZIP = f"{PACKAGE}/09_Supplementary Material 4_extra.zip"


def _bend_workbook():
    """Move one cell of Supplementary Material 2 away from its archived value.

    The tables reach the editor through a build step now, and a build step
    nothing checks is one that will eventually ship an older number than the
    archive holds."""
    from openpyxl import load_workbook
    shutil.copy(SM2_XLSX, SM2_XLSX + ".orig")
    wb = load_workbook(SM2_XLSX)
    ws = wb["S1_signature_retention"]
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, float):
                c.value = c.value + 0.01
                wb.save(SM2_XLSX)
                return
    raise SystemExit("no numeric cell to bend in S1_signature_retention")


def _unbend_workbook():
    os.replace(SM2_XLSX + ".orig", SM2_XLSX)


def _add_archive():
    """Put a compressed file back into the package.

    This is the defect that sent the submission system 259 items: an archive
    uploaded to the journal is unpacked into its members."""
    import zipfile
    with zipfile.ZipFile(_STRAY_ZIP, "w") as z:
        z.writestr("a.txt", "a")
        z.writestr("b.txt", "b")


def _remove_archive():
    if os.path.exists(_STRAY_ZIP):
        os.remove(_STRAY_ZIP)


FILE_MUTATIONS = [
    ("a workbook cell moved off its archived value", _bend_workbook, _unbend_workbook),
    ("a compressed archive left in the package", _add_archive, _remove_archive),
]


def run(path: str) -> str:
    # The checker copy runs from a temporary directory, so the archive cannot be
    # found by walking up from it. IR_ROOT and IR_DOCS say where both are; this
    # is also the exercise that proves those overrides work.
    env = dict(os.environ)
    env["IR_ROOT"] = IR_ROOT
    env["IR_DOCS"] = IR_DOCS
    src = open(CHK, encoding="utf-8").read().replace(
        f'MS = f"{{HOME}}/manuscript_v4.md"', f'MS = "{path}"')
    tmp = path + ".checker.py"
    open(tmp, "w", encoding="utf-8").write(src)
    out = subprocess.run([sys.executable, tmp], capture_output=True, text=True, env=env)
    os.remove(tmp)
    return out.stdout + out.stderr


def main() -> int:
    base = open(SRC, encoding="utf-8").read()
    tmpdir = tempfile.mkdtemp()
    clean = os.path.join(tmpdir, "manuscript_v4.md")
    shutil.copy(SRC, clean)
    print("baseline (unmutated manuscript)")
    out = run(clean)
    ok0 = "ALL CHECKS PASS" in out
    print(("  PASS  " if ok0 else "  FAIL  ") + "the clean manuscript passes every check")
    caught, missed = 0, []
    print("\nmutations")
    for mut in MUTATIONS:
        # A mutation may name a fourth element: a companion file to damage
        # instead of the manuscript. The checker reads the cover letter, the
        # alt text and the submission metadata by absolute path, so those are
        # mutated in place and restored afterwards.
        name, old, new = mut[0], mut[1], mut[2]
        target = mut[3] if len(mut) > 3 else None
        src = open(target, encoding="utf-8").read() if target else base
        if src.count(old) < 1:
            missed.append(f"{name} (anchor not found)")
            print(f"  SKIP  {name}: anchor not found")
            continue
        p = os.path.join(tmpdir, "manuscript_v4.md")
        open(p, "w", encoding="utf-8").write(base if target else
                                             src.replace(old, new))
        if target:
            open(target, "w", encoding="utf-8").write(src.replace(old, new))
        try:
            out = run(p)
        finally:
            if target:
                open(target, "w", encoding="utf-8").write(src)
        got = "CHECK(S) FAILED" in out
        print(("  PASS  " if got else "  FAIL  ") + f"caught: {name}")
        if got:
            caught += 1
        else:
            missed.append(name)
    for name, apply, undo in FILE_MUTATIONS:
        p = os.path.join(tmpdir, "manuscript_v4.md")
        open(p, "w", encoding="utf-8").write(base)
        apply()
        try:
            out = run(p)
        finally:
            undo()
        got = "CHECK(S) FAILED" in out
        print(("  PASS  " if got else "  FAIL  ") + f"caught: {name}")
        if got:
            caught += 1
        else:
            missed.append(name)
    shutil.rmtree(tmpdir, ignore_errors=True)
    total = len(MUTATIONS) + len(FILE_MUTATIONS)
    print(f"\n{caught} of {total} defects caught")
    if missed:
        print("not caught:")
        for m in missed:
            print("  -", m)
    return 0 if (ok0 and not missed) else 1


if __name__ == "__main__":
    sys.exit(main())
